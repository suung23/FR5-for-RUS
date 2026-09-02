#!/usr/bin/env python3
"""검증 리포트 -> Word(.docx).

    python3 report/make_docx.py

본문과 표는 ``make_report`` 한 곳에서만 읽는다. 두 벌로 갈라지면 하나는 반드시
낡는다 — 마크다운과 워드가 다른 숫자를 들고 있는 것이 그 실패다.
"""
from __future__ import annotations

import os
import re
import sys

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_HERE))

from docx import Document                                        # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT                   # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                    # noqa: E402
from docx.oxml import OxmlElement                                # noqa: E402
from docx.oxml.ns import qn                                      # noqa: E402
from docx.shared import Cm, Pt, RGBColor                         # noqa: E402

from report.make_report import REPORT, markdown, tables          # noqa: E402

FIGS = os.path.join(_HERE, "figures")
OUT = os.path.join(_HERE, "force_hold_validation_report.docx")

INK = RGBColor(0x1A, 0x1F, 0x26)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
SUB = RGBColor(0x5A, 0x60, 0x68)
HEAD_FILL = "E8EEF4"
RULE = "C7CCD1"
SANS = "Calibri"

FIG_WIDTH_CM = {"fig1_hold_traces.png": 16.5, "fig2_hold_quality.png": 16.5,
                "fig3_disturbance_response.png": 15.5,
                "fig4_force_exposure.png": 15.5, "fig5_probe_travel.png": 15.5}


def _shade(cell, hexcolor: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def _borders(table) -> None:
    """가로선만 둔다. 세로 격자는 숫자를 읽는 데 방해가 된다."""
    tbl = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8" if edge != "insideH" else "4")
        el.set(qn("w:color"), "1A1F26" if edge != "insideH" else RULE)
        borders.append(el)
    for edge in ("left", "right", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl.append(borders)


def _run(par, text: str, *, bold=False, size=9.5, color=INK):
    r = par.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = SANS
    r.font.color.rgb = color
    return r


def _no_split(row) -> None:
    """행이 쪽 경계에서 쪼개지지 않게 한다."""
    el = OxmlElement("w:cantSplit")
    row._tr.get_or_add_trPr().append(el)


def _widths(table, align, total_cm: float = 16.6) -> None:
    """레이블 열은 넉넉히, 숫자 열은 균등히.

    자동 폭에 맡기면 "Disturbance · 0.5 N" 같은 레이블이 두 줄로 접히고, 표가
    데이터보다 줄바꿈으로 읽힌다.

    셀 폭만 적으면 LibreOffice 는 자동 레이아웃으로 그것을 덮어쓴다. 레이아웃을
    ``fixed`` 로 두고 ``tblGrid`` 까지 같이 적어야 지정한 폭이 살아남는다.
    """
    label = [i for i, a in enumerate(align) if not a.endswith(":")]
    label_cm = 4.4 if label else 0.0
    rest = (total_cm - label_cm * len(label)) / max(1, len(align) - len(label))
    widths = [label_cm if i in label else rest for i in range(len(align))]

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)

    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for cm in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(cm * 567)))          # cm -> twips
        grid.append(col)
    table._tbl.insert(list(table._tbl).index(table._tbl.tblPr) + 1, grid)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])


def add_table(doc, spec) -> None:
    header, rows, align = spec
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)
    for i, name in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        par = cell.paragraphs[0]
        par.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if align[i].endswith(":")
                         else WD_ALIGN_PARAGRAPH.LEFT)
        _run(par, name, bold=True, size=8.5, color=NAVY)
        _shade(cell, HEAD_FILL)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            # 마크다운 강조는 워드에서 굵게 옮긴다 — 별표를 그대로 두면 안 된다.
            bold = text.startswith("**") and text.endswith("**")
            par = cells[i].paragraphs[0]
            par.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if align[i].endswith(":")
                             else WD_ALIGN_PARAGRAPH.LEFT)
            _run(par, text.strip("*"), bold=bold)
    for row in table.rows:
        _no_split(row)
    _widths(table, align)
    doc.add_paragraph()


def add_markdown(doc, text: str, spec: dict) -> None:
    """마크다운 본문을 워드 요소로 옮긴다. 표와 그림은 자리에서 갈아 끼운다."""
    table_order = ["t1", "t2", "t3", "t4"]
    pending_table = iter(table_order)
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()

        if line.startswith("| "):                       # 표 블록 — 통째로 건너뛴다
            while index < len(lines) and lines[index].startswith(("| ", "|-", "|:")):
                index += 1
            add_table(doc, spec[next(pending_table)])
            continue

        if line.startswith("!["):                       # 그림
            name = re.search(r"\(figures/([^)]+)\)", line).group(1)
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(os.path.join(FIGS, name),
                                      width=Cm(FIG_WIDTH_CM.get(name, 15.5)))
            doc.add_paragraph()
            index += 1
            continue

        if line.startswith("# "):
            par = doc.add_paragraph()
            _run(par, line[2:], bold=True, size=19, color=NAVY)
            par.paragraph_format.space_after = Pt(14)
        elif line.startswith("## "):
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(14)
            par.paragraph_format.space_after = Pt(6)
            _run(par, line[3:], bold=True, size=12.5, color=NAVY)
        elif line.startswith("- "):
            # 이어지는 들여쓴 줄은 같은 항목이다. 원문 줄바꿈을 그대로 옮기면
            # 워드에서 한 항목이 여러 줄로 쪼개진 것처럼 보인다.
            item = [line[2:]]
            while index + 1 < len(lines) and lines[index + 1].startswith("  ") \
                    and lines[index + 1].strip():
                index += 1
                item.append(lines[index].strip())
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(4)
            _emphasis(par, " ".join(item), size=10)
        elif line.strip():
            block = [line]
            while index + 1 < len(lines) and lines[index + 1].strip() and \
                    not lines[index + 1].startswith(("|", "!", "#", "- ")):
                index += 1
                block.append(lines[index].rstrip())
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(7)
            joined = " ".join(block)
            # 캡션은 자기 대상과 같은 쪽에 있어야 한다. 떨어지면 독자가 무엇의
            # 캡션인지 앞 쪽으로 되돌아가 확인해야 한다.
            if joined.startswith(("**Figure", "**Table")):
                par.paragraph_format.keep_with_next = True
                par.paragraph_format.space_before = Pt(6)
            _emphasis(par, joined, size=10)
        index += 1


def _emphasis(par, text: str, *, size: float) -> None:
    """``**굵게**`` 와 ``` `코드` ``` 를 워드 서식으로."""
    for piece in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not piece:
            continue
        if piece.startswith("**"):
            _run(par, piece.strip("*"), bold=True, size=size)
        elif piece.startswith("`"):
            r = _run(par, piece.strip("`"), size=size - 0.5, color=SUB)
            r.font.name = "Consolas"
        else:
            _run(par, piece, size=size)


def main() -> int:
    spec = tables()
    text = REPORT.format(**{k: markdown(v) for k, v in spec.items()})

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for attr, value in (("left_margin", 2.2), ("right_margin", 2.2),
                        ("top_margin", 2.0), ("bottom_margin", 2.0)):
        setattr(section, attr, Cm(value))
    style = doc.styles["Normal"]
    style.font.name = SANS
    style.font.size = Pt(10)

    add_markdown(doc, text, spec)
    doc.save(OUT)
    print(os.path.relpath(OUT, os.path.dirname(_HERE)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
