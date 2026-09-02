#!/usr/bin/env python3
"""Sonologger 원고(.docx) 조립 — 본문·그림·참고문헌.

    python3 docs/build_sonologger_manuscript.py
    soffice --headless --convert-to pdf Sonologger_Manuscript.docx

본문은 이 파일 하나에만 있다. 그림은 리포지토리의 기존 자산을 그대로 넣는다
(재생성하지 않는다). 근거 추적은 MANUSCRIPT_EVIDENCE_AND_OPEN_ITEMS.md 에 있다.
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIGDIR = os.path.join(REPO, "imu_bench", "qc_track", "report", "figures")
OUT = os.path.join(REPO, "Sonologger_Manuscript.docx")

BODY_PT = 9.2
SERIF = "Times New Roman"

#: 폭 [cm]. 원본은 슬라이드 크기(16.6 in)로 그려져 있어 A4 본문 폭으로 줄이면
#: 가장 작은 주석이 4 pt 근처가 된다. 그래서 판독이 걸린 그림은 본문 폭을 꽉 채운다.
FIGS = {
    1: ("fig_sonologger_acquisition.png", 17.4),
    2: ("fig_6axis_rotation_translation.png", 15.0),
    3: ("fig_learning_to_control.png", 16.0),
}


# ----------------------------------------------------------------- 서식 도구
def style_doc(doc):
    st = doc.styles["Normal"]
    st.font.name = SERIF
    st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0

    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin = s.bottom_margin = Cm(1.7)
        s.left_margin = s.right_margin = Cm(1.7)
        s.header_distance = s.footer_distance = Cm(1.0)


def page_numbers(doc):
    """가운데 정렬 페이지 번호 — 필드 코드로 넣는다."""
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.font.name, r.font.size = SERIF, Pt(9)
    for txt, kind in (("begin", "fldChar"), ("PAGE", "instrText"), ("end", "fldChar")):
        el = OxmlElement("w:" + kind)
        if kind == "fldChar":
            el.set(qn("w:fldCharType"), txt)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        r._r.append(el)


def para(doc, text, *, size=BODY_PT, bold=False, italic=False, align="just",
         before=0, after=0, indent=0.0, color=None, first_line=None):
    p = doc.add_paragraph()
    p.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.left_indent = Cm(indent)
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    rich(p, text, size=size, bold=bold, italic=italic, color=color)
    return p


def rich(p, text, *, size=BODY_PT, bold=False, italic=False, color=None):
    """'**굵게**', '*기울임*', '~아래첨자~', '^위첨자^' 를 런으로 나눈다."""
    import re
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|~[^~]+?~|\^[^^]+?\^)", text):
        if not tok:
            continue
        b, i, sub, sup = bold, italic, False, False
        if tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*"):
            tok, i = tok[1:-1], True
        elif tok.startswith("~") and tok.endswith("~"):
            tok, sub = tok[1:-1], True
        elif tok.startswith("^") and tok.endswith("^"):
            tok, sup = tok[1:-1], True
        r = p.add_run(tok)
        r.font.name, r.font.size = SERIF, Pt(size)
        r.font.bold, r.font.italic = b, i
        r.font.subscript, r.font.superscript = sub, sup
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def head(doc, text, level=1):
    size = {1: 10.2, 2: 9.5}[level]
    p = para(doc, text, size=size, bold=True, align="left",
             before=5.5 if level == 1 else 4, after=1.5)
    p.paragraph_format.keep_with_next = True      # 쪽 끝에 제목만 남지 않게


def figure(doc, n, caption):
    name, width_cm = FIGS[n]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.keep_with_next = True      # 그림과 설명을 갈라놓지 않는다
    p.add_run().add_picture(os.path.join(FIGDIR, name), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    c.paragraph_format.space_after = Pt(5)
    rich(c, "**Figure %d.** %s" % (n, caption), size=8.0)


def equation(doc, text, tag):
    """가운데 정렬한 식 + 오른쪽 끝의 식 번호 (탭 정지로 고정)."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(3), Pt(3)
    pf.tab_stops.add_tab_stop(Cm(8.8), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(17.6), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    rich(p, text, size=BODY_PT, italic=True)
    p.add_run("\t")
    rich(p, "(%s)" % tag, size=BODY_PT)


def bullets(doc, items, size=BODY_PT):
    for it in items:
        para(doc, it, size=size, indent=0.45, first_line=-0.45, after=1)


# ----------------------------------------------------------------------- 본문
def build():
    doc = Document()
    style_doc(doc)
    page_numbers(doc)

    # ---- 제목 ----------------------------------------------------------
    para(doc, "Sonologger: A Synchronized Ultrasound–IMU Platform for "
              "Robot-Free Acquisition of Robotic Ultrasound Learning Data",
         size=14, bold=True, align="center", after=6)

    para(doc, "[TO BE CONFIRMED: Given Name Surname]^1^, "
              "[TO BE CONFIRMED: Given Name Surname]^1^, "
              "[TO BE CONFIRMED: Given Name Surname]^2^, "
              "[TO BE CONFIRMED: Corresponding Author]^1,^*",
         size=10, align="center", after=3)
    para(doc, "^1^ [TO BE CONFIRMED: Department / Institution, City, Republic of Korea]",
         size=8.6, align="center", after=0)
    para(doc, "^2^ [TO BE CONFIRMED: Department of Urology, Institution, City, "
              "Republic of Korea]", size=8.6, align="center", after=0)
    para(doc, "* Correspondence: [TO BE CONFIRMED: e-mail address]",
         size=8.6, align="center", after=7)

    # ---- 초록 ----------------------------------------------------------
    head(doc, "Abstract", 1)
    para(doc,
         "**Objectives.** Learning-based robotic ultrasound requires large collections of "
         "expert scanning demonstrations, yet the usual routes to that data — mounting the "
         "probe on a robot, instrumenting the room with an external tracker, or working in a "
         "dedicated laboratory — restrict where and how often recordings can be made. We "
         "present Sonologger, a lightweight platform that records synchronized ultrasound "
         "video and probe motion while a clinician scans by hand, with no robot in the loop. "
         "**Methods.** A printed clamp shell fits a GE 4C-RS convex probe and carries a "
         "BNO085 inertial measurement unit (IMU). The display output of a GE Voluson P6 "
         "scanner is captured over HDMI while the IMU streams over USB serial; the host "
         "timestamps both streams with a single PC Unix clock at reception, so image frames "
         "and motion records join on one key. Orientation is estimated with a "
         "magnetometer-free Madgwick filter and expressed relative to a stationary reference "
         "pose. **Results.** Bench sessions recorded 776 ultrasound frames over 96.9 s "
         "(8.00 fps) alongside 24,471 IMU records (252 Hz) on the shared clock. Against robot "
         "forward-kinematics ground truth during teleoperated probing-like motion, "
         "latency-corrected angular root-mean-square error was 1.14°, 0.73° and 1.95° for "
         "tilt x, tilt y and axial rotation, with 0.23°/min heading drift. "
         "**Conclusions.** Sonologger lowers the barrier to collecting scanning "
         "demonstrations. It records relative orientation, not translation or complete pose; "
         "downstream policy learning remains future work.", after=4)

    para(doc, "**Keywords:** robotic ultrasound; learning from demonstration; inertial "
              "measurement unit; data acquisition; time synchronization; bladder ultrasound",
         after=2)

    # ---- 1 서론 --------------------------------------------------------
    head(doc, "1. Introduction", 1)
    para(doc,
         "Diagnostic ultrasound is inexpensive, real-time and free of ionizing radiation, but "
         "examination quality depends heavily on how the operator holds and moves the probe. "
         "Robotic ultrasound systems have been proposed to make probe placement and contact "
         "more reproducible [1–4]. Our target workflow is narrower than a full examination: "
         "during holmium laser enucleation of the prostate (HoLEP), an assisting clinician "
         "repeatedly re-establishes a transabdominal view of the bladder lumen so that the "
         "relationship between the bladder wall and the morcellator stays visible. The task is "
         "repetitive, constrained, and performed under time pressure by personnel whose "
         "availability is itself a constraint [5,6] — the profile of a task worth "
         "standardizing.")
    para(doc,
         "Contemporary robot manipulation policies are learned from demonstrations rather than "
         "hand-designed [7–10]. Transferring that recipe to ultrasound requires many paired "
         "sequences of image and probe motion produced by people who already scan well, and "
         "this is where the field is constrained. Such data are normally acquired with the "
         "probe held by the robot, with an optical or electromagnetic tracker observing the "
         "probe, or in a laboratory arranged for the purpose. Each option changes what is "
         "recorded: a robot-held probe removes the human behaviour to be imitated, external "
         "trackers need line of sight and per-session calibration, and a laboratory limits the "
         "case mix. In manipulation, the Universal Manipulation Interface showed that this "
         "trade can be avoided by instrumenting a hand-held tool instead of the robot [10].")
    para(doc,
         "Freehand three-dimensional ultrasound has combined a tracked probe with B-mode video "
         "for decades [11,12], usually with electromagnetic or optical tracking and more "
         "recently by estimating inter-frame motion from the images themselves [13]. Inertial "
         "sensing is cheap, unobtrusive and compatible with a clinical room, but it measures "
         "orientation; recovering position from it requires double integration, which is "
         "dominated by bias growth unless the motion is periodically brought to rest [14,15]. "
         "The defensible claim for an IMU-only rig is therefore orientation, not pose, and the "
         "design presented here is built around that constraint rather than against it.")
    para(doc, "This paper contributes:", before=2, after=1)
    bullets(doc, [
        "•\ta robot-free acquisition platform pairing clinical ultrasound display video with "
        "on-probe inertial data on a single host clock, using an unmodified scanner and probe;",
        "•\ta reference-pose convention with a gated stationary calibration, making relative "
        "orientation reproducible across sessions and operators;",
        "•\ta characterization of the orientation channel against robot forward-kinematics "
        "ground truth, including the evidence for disabling the magnetometer; and",
        "•\tan explicit account of what the platform does not measure, with the session format "
        "by which recordings are intended to become learning data.",
    ])
    para(doc,
         "No robotic learning model has been trained, validated or deployed on Sonologger "
         "data; robot learning is the intended downstream application and is identified as "
         "such throughout.", before=1.5)

    # ---- 2 재료와 방법 --------------------------------------------------
    head(doc, "2. Materials and Methods", 1)

    head(doc, "2.1. System overview", 2)
    para(doc,
         "Sonologger consists of an instrumented probe, two independent acquisition paths, and "
         "one host computer that timestamps both (Figure 1). The image path carries what the "
         "operator sees; the pose path carries how the probe is oriented while they see it. "
         "Neither path requires a robot.")

    head(doc, "2.2. Instrumented probe", 2)
    para(doc,
         "The probe housing was surveyed dimensionally, reconstructed as a CAD solid, and used "
         "to design a hinged split shell that clamps around the existing housing and carries a "
         "sensor bracket (Figure 1a). The printed shell is added before a session and removed "
         "afterwards; neither the probe nor the scanner is modified, and no internal signal is "
         "taken from either. The IMU is rigidly fixed with respect to the transducer array, so "
         "the rotation it reports is the rotation of the imaging plane. We use a probe frame "
         "in which +z is the beam axis directed into tissue, +x lies along the array inside "
         "the image plane, and +y is the elevational direction normal to it; rotations about "
         "these axes are termed tilt (out-of-plane), rock (in-plane) and axial rotation.")

    head(doc, "2.3. Ultrasound acquisition", 2)
    para(doc,
         "The GE 4C-RS probe is connected to a GE Voluson P6 scanner and used normally. The "
         "scanner's display output is taken over HDMI into a USB video capture device, and the "
         "host timestamps each captured display frame on arrival. Capturing the display rather "
         "than beamformed data keeps the platform scanner-agnostic and requires no access to "
         "the machine's internals, but it also fixes two properties of the dataset: the "
         "recorded image is exactly what the operator saw, including overlays and the "
         "machine's own post-processing, and the effective frame rate is that of the capture "
         "chain — approximately 8 fps here — rather than the scanner's internal rate. Frames "
         "are written as raw bytes without reorientation or scan conversion, and a per-frame "
         "index records the host timestamp, sequence number and byte offset into the frame "
         "store.")

    figure(doc, 1,
           "Sonologger acquisition. (a) Build of the instrumented probe: dimensional survey of "
           "the GE 4C-RS convex array, CAD reconstruction, printed hinged split clamp shell "
           "with sensor bracket, and the BNO085 breakout on that bracket. (b) Acquisition "
           "architecture: the inertial stream reaches the host over USB serial while the "
           "ultrasound stream leaves the scanner as display video, is digitized by an HDMI "
           "capture device, and arrives as captured display frames; both are stamped with the "
           "host clock on arrival and stored in one session directory whose join key is "
           "pc_unix. (c) The resulting common time base. The arm in (b) is the bench fixture "
           "used in Section 3, not part of the acquisition path.")

    head(doc, "2.4. Inertial acquisition and orientation estimation", 2)
    para(doc,
         "The IMU is a BNO085 on a microcontroller board. Firmware requests calibrated "
         "accelerometer and gyroscope reports at 200 Hz each and calibrated magnetometer and "
         "on-chip rotation-vector reports at 100 Hz; records are framed with an 8-bit cyclic "
         "redundancy check, carry a device timestamp in microseconds, and reach the host over "
         "USB serial at 115,200 baud. Host-side attitude is estimated with a Madgwick "
         "attitude-and-heading reference filter [16] in a north-west-up earth frame; with no "
         "magnetometer vector supplied it uses the six-axis update, which is the configuration "
         "the platform runs, for reasons measured in Section 3.3. Each logged row nevertheless "
         "retains the raw and calibrated magnetometer vectors, the host and on-chip "
         "quaternions and the sensor's calibration status, so this choice can be revisited "
         "offline without re-acquiring data.")

    head(doc, "2.5. Reference-pose calibration", 2)
    para(doc,
         "Filtered attitude is referenced to the earth frame and therefore bears no fixed "
         "relation to the patient, the bed or the scanner. Every session begins with a short "
         "stationary hold defining a reference from which subsequent orientation is read:")
    equation(doc, "R~rel~(t) = R~0~^T^ R(t),        q~rel~(t) = q~0~^−1^ ⊗ q(t)", "1")
    para(doc,
         "Here R(t) ∈ SO(3) is the estimated rotation from the sensor frame to the earth frame "
         "at time t and q(t) the corresponding unit quaternion; R~0~ and q~0~ are their values "
         "during the reference hold; R~0~^T^ denotes matrix transpose, q~0~^−1^ the conjugate "
         "of the unit quaternion q~0~, and ⊗ quaternion multiplication. R~rel~(t) and "
         "q~rel~(t) are two representations of the same quantity: the rotation accumulated "
         "since the reference. The hold is accepted only if gyroscope standard deviation and "
         "mean are below 0.005 rad/s and accelerometer standard deviation is below 0.15 m/s²; "
         "a failing hold is refused and reported rather than silently recorded, because a "
         "reference captured during motion is indistinguishable from a valid one in the log. "
         "The same hold estimates the earth-frame accelerometer bias, whose magnitude is "
         "reported as a measured gravity value and a scale error.")
    para(doc,
         "Equation (1) yields relative orientation only; it does not provide probe position. A "
         "pose state T~probe~(t) ∈ SE(3), should a later control stage require one, would have "
         "to come from an additional modality — robot kinematics, an external tracker, or "
         "image-based motion estimation — and is named here only as a future or externally "
         "augmented representation. The present IMU-only configuration does not produce it.")

    head(doc, "2.6. Time synchronization and frame–motion association", 2)
    para(doc,
         "The two streams reach the host by different physical routes and at different rates. "
         "They are placed on one axis by the simplest available means: the host reads its own "
         "Unix clock as each record completes and writes that value, pc_unix, into both logs. "
         "Association is a join on that key — for each ultrasound frame, the inertial records "
         "nearest in pc_unix. Two properties of the scheme are stated in the session metadata "
         "rather than left implicit. First, it is arrival-time alignment: residual drift "
         "between host and device clocks is measurable from the device timestamp and "
         "correctable after the fact. Second, the fixed end-to-end delay of the ultrasound "
         "path — beamforming, display refresh, capture, USB transport — is not removed by "
         "arrival stamping and has not yet been measured [TO BE CONFIRMED: ultrasound "
         "display-to-host latency]; being a fixed delay rather than jitter, it appears as a "
         "constant image-to-motion offset and can be compensated once measured against a "
         "common physical event.")

    head(doc, "2.7. Session organization", 2)
    para(doc,
         "A session is one directory holding the per-sample inertial log as CSV with a sidecar "
         "metadata file, the ultrasound frames as one concatenated binary store, a per-frame "
         "index (host timestamp, sequence number, frame identifier, byte offset), and session "
         "metadata naming the shared time axis, the join key and the known caveats. The "
         "inertial log carries raw and calibrated inertial and magnetic vectors, host and "
         "on-chip attitude, the reference-relative orientation of Equation (1), and per-sample "
         "calibration status, so every estimator choice in this paper can be reversed offline.")

    head(doc, "2.8. Intended conversion to robotic learning data", 2)
    para(doc,
         "Figure 3 shows how recorded sessions are intended to become training data and how a "
         "resulting policy would be embedded in a control stack. Sessions would be segmented "
         "into brackets bounded by stillness, so that net motion within a bracket is "
         "observable even where per-step velocity labels are not, and each bracket treated as "
         "one action chunk paired with a short window of raw frames and the recorded state. "
         "This pipeline is a design: it has not been trained on Sonologger data, and no "
         "component of the control stack in Figure 3b has been evaluated with a policy learned "
         "from this platform.")

    figure(doc, 3,
           "Intended downstream use of Sonologger sessions — a planned workflow, not a result. "
           "(a) Offline: sessions joined on the host clock, segmented into stillness-bounded "
           "brackets, and converted into observation–action chunks for policy training. "
           "(b) Online: the control stack such a policy is intended to drive. No policy has "
           "been trained on Sonologger data and no component of (b) has been evaluated with "
           "one; the quantities shown are design targets from the system specification.")

    # ---- 3 시연 --------------------------------------------------------
    head(doc, "3. System Demonstration and Preliminary Evaluation", 1)
    head(doc, "3.1. Synchronized acquisition", 2)
    para(doc,
         "This section reports what was measured; it evaluates neither a learned policy, nor "
         "image quality, nor clinical performance, none of which was attempted. Five bench "
         "sessions were recorded and three contain both streams. In the longest, 776 "
         "ultrasound frames were captured over 96.9 s — 8.00 frames per second, median "
         "inter-frame interval 125.9 ms — alongside 24,471 inertial records over 97.0 s, a "
         "mean logged rate of 252 Hz. Every frame carries a host timestamp on the same axis as "
         "every inertial row, so the frame-to-motion join is direct and needs no interpolation "
         "of one stream onto the other's nominal grid. No dropped frame or reconnection event "
         "was recorded. These particular sessions were captured with the host collector "
         "reading 256 × 256 frames from a research image front end; the host timestamping and "
         "join logic are identical for either front end, but no session recorded through the "
         "HDMI display path of Figure 1b is yet stored "
         "[TO BE CONFIRMED: end-to-end sessions through the HDMI capture path].")

    head(doc, "3.2. Orientation fidelity against robot ground truth", 2)
    para(doc,
         "To characterize the orientation channel against an independent reference, the "
         "instrumented probe was mounted on the flange of a six-degree-of-freedom "
         "collaborative arm and the arm teleoperated through a probing-like rhythm — position, "
         "press, hold at least 1.5 s while looking at the image, move on. Robot forward "
         "kinematics served as ground truth at approximately 144 Hz while the IMU logged at "
         "approximately 253 Hz. Pairing between the two logs was complete, with −25 ppm clock "
         "skew and a pairing residual 95th percentile of 4.9 ms, so the attitude error below "
         "is not a timing artefact.")

    para(doc,
         "Over an 83.8-s block the effective latency of the orientation channel was −22.8, "
         "−21.1 and −6.6 ms for tilt x, tilt y and axial rotation, with cross-correlations of "
         "0.997, 0.999 and 1.000 against ground truth (Figure 2a–b). After latency correction, "
         "angular RMS errors were 1.14°, 0.73° and 1.95° for the three channels (95th "
         "percentiles 2.23°, 1.29°, 2.57°; Figure 2c). Tilt magnitude — the quantity setting "
         "the inclination of the imaging plane — had an RMS error of 0.42°. Treating the "
         "rotation as a whole, the geodesic attitude error was 2.50° RMS. Heading drift over "
         "the block was 0.23°/min and the static alignment residual across seven stationary "
         "poses 0.81° RMS.")
    para(doc,
         "Against acceptance criteria fixed before the run, alignment residual, latency, "
         "cross-correlation, heading drift and displacement error floor all passed, while the "
         "geodesic attitude RMS of 2.50° exceeded the 2.0° limit set for it. We report that "
         "outcome as recorded rather than adjusting the threshold; the per-channel errors that "
         "matter most for image-plane orientation are below 2°, and the discrepancy is "
         "concentrated in the axial-rotation channel.")

    head(doc, "3.3. Why the magnetometer is disabled", 2)
    para(doc,
         "The same raw recordings were re-processed with and without magnetic information. "
         "The on-chip nine-axis rotation vector was worse on every rotational channel — 11.23° "
         "against 1.95° in axial rotation (×5.8), 6.17° against 1.14° in tilt x (×5.4), 3.98° "
         "against 0.73° in tilt y (×5.4) and 0.70° against 0.42° in tilt magnitude (×1.7) — "
         "and its static alignment residual was 6.68° against 0.81° (×8.2), its heading drift "
         "8.29°/min against 0.23°/min (×36.1).")

    para(doc,
         "The cause was localized rather than assumed. A test searching for stationary pairs "
         "of equal attitude but different position — pairs for which any sensor-fixed "
         "distortion model must predict an identical field — found no usable pair across three "
         "blocks, and a fitted model left residuals up to 9.6 times the sensor noise with "
         "direction errors up to 56.8° in one block. The field varied with position, not "
         "merely with sensor attitude, which is the regime in which magnetometer calibration "
         "cannot help [17]. We therefore run with the magnetometer disabled and accept that "
         "absolute heading is then unobservable — precisely why orientation is reported "
         "relative to a per-session reference pose (Equation 1): the platform claims rotation "
         "since a known stationary moment, not a compass bearing.")

    head(doc, "3.4. What the platform does not measure", 2)
    para(doc,
         "Applying zero-velocity-updated double integration to stationary segments gives an "
         "error floor of 0.25, 0.67, 1.97 and 6.32 mm for integration windows of 0.5, 1, 2 and "
         "5 s; without bias removal and zero-velocity updates the same windows give 12.4, "
         "49.5, 197.8 and 1236.5 mm (Figure 2e). On 21 actual moving segments of 20–115 mm "
         "travel, displacement error was one to two orders of magnitude above that floor and "
         "grew with segment duration (Figure 2d, f), consistent with an independent bench "
         "measurement of the same sensor. Sonologger therefore records probe orientation; it "
         "does not record probe displacement, and no such claim is made.")

    figure(doc, 2,
           "Orientation and displacement of the magnetometer-free channel against robot "
           "forward-kinematics ground truth over an 83.8-s teleoperated probing block. "
           "(a) Time series per rotational channel, annotated with effective latency, "
           "cross-correlation and latency-corrected root-mean-square (RMS) error. (b) Detail "
           "of the shaded window after latency correction. (c) Distribution of the "
           "latency-corrected attitude error. (d) Displacement error per zero-velocity-bounded "
           "segment against travel distance (n = 21). (e) Displacement error floor on "
           "stationary segments, for four integration windows and three correction stages. "
           "(f) That floor against the error actually incurred while moving. Panels (d)–(f) "
           "are the basis for the statement that the platform records orientation and not "
           "displacement. The motion is robot-actuated, not a clinician\'s hand.")

    # ---- 4 논의 --------------------------------------------------------

    head(doc, "4. Discussion", 1)
    para(doc,
         "Nothing in the acquisition path is specific to a robot or to a laboratory. The clamp "
         "shell is added to an existing probe and removed after the session, the scanner is "
         "used unmodified and observed only through its display output, and the sensor is a "
         "USB device. Demonstrations can therefore be recorded where the procedure actually "
         "takes place and by the clinician who is otherwise the only source of the skill. In "
         "the intended clinical direction — bladder ultrasound monitoring during HoLEP "
         "morcellation — recording is passive and non-interventional: the existing surgical "
         "workflow is preserved, no additional manipulation or examination is requested, and "
         "no robot, force–torque sensor or automatic contact-force control is applied to a "
         "research subject. No clinical data have been collected and the ethics application is "
         "in preparation.")

    para(doc,
         "Relative to robotic ultrasound systems [1–4], Sonologger is not a robot and performs "
         "no control; it removes the robot from the data-collection stage only, and the broader "
         "research programme it belongs to does involve one. Relative to freehand "
         "three-dimensional ultrasound [11–13], we attempt no volume reconstruction and claim "
         "none of the position accuracy that would require. Relative to inertial motion capture "
         "[14–16,18], the contribution is not the estimator — a standard Madgwick "
         "implementation — but its pairing with clinical display video on one host clock and a "
         "measured account of what that pairing does and does not resolve. Relative to "
         "robot-free demonstration collection in manipulation [10], the argument is transferred "
         "to a domain where the action is largely a rotation of an imaging plane.")
    para(doc, "Several limitations bound what this platform can currently support.", before=2, after=1)
    bullets(doc, [
        "•\t**Display-capture rate.** At approximately 8 fps, rapid probe motion is "
        "undersampled relative to the scanner's own frame rate, and the recorded image is the "
        "rendered display, so machine post-processing and overlays are baked in and beamformed "
        "data are unavailable.",
        "•\t**Arrival-time uncertainty.** Inter-stream clock skew is small and correctable, "
        "but the fixed ultrasound path delay is neither removed by arrival stamping nor yet "
        "measured, so frame-to-motion association carries an unknown constant offset.",
        "•\t**Orientation limits.** The reported errors come from robot-actuated motion over "
        "blocks of roughly 84 s, not from a clinician's hand over a whole procedure. With the "
        "magnetometer disabled only drift relative to the session reference is meaningful, and "
        "the 0.23°/min figure was measured in one environment.",
        "•\t**No translation and no complete pose** (Section 3.4). Any application requiring "
        "T~probe~(t) ∈ SE(3) must add a modality.",
        "•\t**No downstream validation.** No policy has been trained on Sonologger data, and "
        "no segmentation or image-quality result is claimed.",
        "•\t**Bench data only.** There is no operator study, no case series and no clinical "
        "dataset.",
    ])
    para(doc,
         "Next steps follow from that list: measure the ultrasound end-to-end latency against "
         "a common physical event so the constant offset can be removed; record freehand "
         "sessions with a clinician under an approved protocol; and determine empirically "
         "whether relative orientation paired with the image is sufficient supervision for the "
         "target task, or whether the translational component must come from another "
         "modality.", before=1.5)

    # ---- 5 결론 --------------------------------------------------------
    head(doc, "5. Conclusion", 1)
    para(doc,
         "Sonologger records clinical ultrasound video and on-probe inertial motion on a "
         "single host clock, using an unmodified scanner and probe and requiring no robot, no "
         "external tracker and no dedicated room. Bench sessions demonstrate stable "
         "simultaneous acquisition at approximately 8 frames and 252 inertial records per "
         "second on a shared time axis, and comparison against robot forward kinematics places "
         "the latency-corrected angular error of the magnetometer-free channel between 0.4° "
         "and 2.0° on the axes that determine image-plane orientation. The platform supplies "
         "relative orientation from a stationary reference, not translation and not complete "
         "pose. Whether such recordings suffice to train a robotic ultrasound policy is an "
         "open question and the subject of future work.")

    # ---- 선언 ----------------------------------------------------------
    head(doc, "Declarations", 1)
    para(doc,
         "**Ethics approval.** No human-subject data were collected for this work; all "
         "recordings reported here are bench sessions. The planned clinical extension is a "
         "non-interventional observational study and will not begin before institutional "
         "review board approval [TO BE CONFIRMED: IRB name and approval number]. "
         "**Funding.** [TO BE CONFIRMED: funding source and grant number]. "
         "**Conflicts of interest.** [TO BE CONFIRMED: declaration]. "
         "**Data and code availability.** [TO BE CONFIRMED: repository URL and access "
         "policy]. **Author contributions.** [TO BE CONFIRMED: CRediT statement].",
         size=8.8)

    # ---- 참고문헌 ------------------------------------------------------
    head(doc, "References", 1)
    refs = [
        "Priester AM, Natarajan S, Culjat MO. Robotic ultrasound systems in medicine. IEEE "
        "Trans Ultrason Ferroelectr Freq Control. 2013;60(3):507–523.",
        "von Haxthausen F, Böttger S, Wulff D, Hagenah J, García-Vázquez V, Ipsen S. Medical "
        "robots for ultrasound imaging: current systems and future trends. Curr Robot Rep. "
        "2021;2:55–71.",
        "Li K, Xu Y, Meng MQ-H. An overview of systems and techniques for autonomous robotic "
        "ultrasound acquisitions. IEEE Trans Med Robot Bionics. 2021;3(2):510–524.",
        "Jiang Z, Salcudean SE, Navab N. Robotic ultrasound imaging: state-of-the-art and "
        "future perspectives. Med Image Anal. 2023;89:102878.",
        "Robles et al. [TO BE CONFIRMED: full author list with initials, title, volume and "
        "pages]. J Endourol. 2025.",
        "Tzou et al. [TO BE CONFIRMED: full author list with initials, title, volume and "
        "pages]. Urology. 2020.",
        "Argall BD, Chernova S, Veloso M, Browning B. A survey of robot learning from "
        "demonstration. Robot Auton Syst. 2009;57(5):469–483.",
        "Ravichandar H, Polydoros AS, Chernova S, Billard A. Recent advances in robot learning "
        "from demonstration. Annu Rev Control Robot Auton Syst. 2020;3:297–330.",
        "Zhao TZ, Kumar V, Levine S, Finn C. Learning fine-grained bimanual manipulation with "
        "low-cost hardware. In: Proceedings of Robotics: Science and Systems (RSS); 2023.",
        "Chi C, Xu Z, Pan C, Cousineau E, Burchfiel B, Feng S, Tedrake R, Song S. Universal "
        "Manipulation Interface: in-the-wild robot teaching without in-the-wild robots. In: "
        "Proceedings of Robotics: Science and Systems (RSS); 2024.",
        "Prager RW, Gee A, Berman L. Stradx: real-time acquisition and visualization of "
        "freehand three-dimensional ultrasound. Med Image Anal. 1999;3(2):129–140.",
        "Housden RJ, Treece GM, Gee AH, Prager RW. Sensorless reconstruction of unconstrained "
        "freehand 3D ultrasound data. Ultrasound Med Biol. 2007;33(3):408–419.",
        "Prevost R, Salehi M, Jagoda S, Kumar N, Sprung J, Ladikos A, Bauer R, Zettinig O, "
        "Wein W. 3D freehand ultrasound without external tracking using deep learning. Med "
        "Image Anal. 2018;48:187–202.",
        "Foxlin E. Pedestrian tracking with shoe-mounted inertial sensors. IEEE Comput Graph "
        "Appl. 2005;25(6):38–46.",
        "Skog I, Händel P, Nilsson J-O, Rantakokko J. Zero-velocity detection — an algorithm "
        "evaluation. IEEE Trans Biomed Eng. 2010;57(11):2657–2666.",
        "Madgwick SOH, Harrison AJL, Vaidyanathan R. Estimation of IMU and MARG orientation "
        "using a gradient descent algorithm. In: Proceedings of the IEEE International "
        "Conference on Rehabilitation Robotics (ICORR); 2011. p. 1–7.",
        "de Vries WHK, Veeger HEJ, Baten CTM, van der Helm FCT. Magnetic distortion in motion "
        "labs, implications for validating inertial magnetic sensors. Gait Posture. "
        "2009;29(4):535–541.",
        "Sabatini AM. Estimating three-dimensional orientation of human body parts by "
        "inertial/magnetic sensing. Sensors. 2011;11(2):1489–1525.",
    ]
    for i, r in enumerate(refs, 1):
        para(doc, "%d.\t%s" % (i, r), size=7.9, indent=0.5, first_line=-0.5, after=0.5)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
